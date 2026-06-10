"""只更新修正版PPT的技术难点页(P17), 不动其他页"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

path = os.path.expanduser(r"C:\Users\ASUS\OneDrive\Desktop\desktopdog\修正版\桌面智能灵动岛_答辩PPT（修正版）.pptx")
prs = Presentation(path)

BG = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0x00, 0xD2, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0xCC, 0x71)
ORANGE = RGBColor(0xFF, 0x9F, 0x43)
RED = RGBColor(0xFF, 0x6B, 0x6B)
PURPLE = RGBColor(0xA8, 0x55, 0xF7)
GRAY = RGBColor(0xAA, 0xAA, 0xAA)
CARD = RGBColor(0x25, 0x25, 0x3D)

# slide index 16 = P17
sl = prs.slides[16]

# clear all existing shapes except background
for shape in list(sl.shapes):
    sp = shape._element
    sp.getparent().remove(sp)

def tb(l, t, w, h, txt, sz=18, c=WHITE, bold=False, align=PP_ALIGN.LEFT):
    b = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    f = b.text_frame; f.word_wrap = True
    p = f.paragraphs[0]; p.text = txt
    p.font.size = Pt(sz); p.font.color.rgb = c; p.font.bold = bold
    p.font.name = 'Microsoft YaHei'; p.alignment = align
    return f

def card(l, t, w, h, tt, bt, tc=ACCENT, bc=WHITE):
    c = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    c.fill.solid(); c.fill.fore_color.rgb = CARD; c.line.fill.background()
    tb(l+0.15, t+0.10, w-0.3, 0.4, tt, sz=15, c=tc, bold=True)
    tb(l+0.15, t+0.50, w-0.3, h-0.6, bt, sz=10.5, c=bc)

# Title
tb(0.8, 0.4, 11.7, 1.0, "关键技术难点与解决方案", sz=36, c=ACCENT, bold=True)
l = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.15), Inches(2.0), Inches(0.04))
l.fill.solid(); l.fill.fore_color.rgb = ACCENT; l.line.fill.background()

# ── 4 cards, 2x2 grid ──
# Row 1
card(0.6, 1.4, 6.0, 2.8,
    "难点一：GPIO 引脚资源极度紧缺",
    "问题: ESP32-C3 仅 11 个可用 GPIO，需同时驱动 OLED/麦克风/功放/按键/电量检测\n\n"
    "尝试方案: 74HC138 3-8译码器扩展按键 (4脚→6键) → 焊接问题致芯片损坏，放弃\n"
    "GPIO0 物理损坏 (静电/过温) → 损失 1 键，改为电池 ADC 分压检测\n\n"
    "退化过程: 6键+LED → HC138 损坏 → 5键无LED → GPIO0 损坏 → 最终 4 键\n"
    "4 键替代方案: D 键扛 3 事件 (短按切模式/双击配对/长按进设置)\n"
    "A 键按模式变换角色 (CHAT-ROBOT 长按录音 / COUNTDOWN 短按-5min)\n"
    "全用 INPUT_PULLUP 省外部电阻 · GPIO2 因 Strapping 特殊处理为 INPUT\n\n"
    "代码: config.h KEY_PIN_A=2,B=4,C=3,D=1 · button_hal.cpp 1kHz 轮询+8次去抖",
    RED)

card(6.9, 1.4, 5.5, 2.8,
    "难点二：语音全链路 (两代架构演进)",
    "第一代 (ESP32直连智谱AI→OLED):\n"
    "  • ai_service.cpp 限制「5 英文单词」→ 回复无信息量，方案失败\n"
    "  • 为什么限制 5 词？OLED 128×64 内容区仅 48px 高，中文一行 8 字\n\n"
    "第二代 (ESP32→MQTT→Web DeepSeek Agent→MQTT→ESP32):\n"
    "  • agent.js: think→tool_call→observe 多轮推理 + 8 个 Function Calling\n"
    "  • 回复质量好但延迟长: 2s 录音 + 7-12s 处理，用户松手等太久\n\n"
    "关键技术创新 — 流式 base64 编码 (asr_service.cpp):\n"
    "  • ESP32-C3 仅 400KB RAM → 无法一次构建 WAV+base64 (128KB)\n"
    "  • streamBase64Concat(): WAV头+PCM→480B 分块边编码边发送\n"
    "  • WAV头与PCM必须作为连续流 → 避免错误 (=) 填充符致 API 失败\n\n"
    "遗留: 15s WEB_AI_TIMEOUT 太短，Agent 多轮 tool_call 偶尔超时",
    ORANGE)

# Row 2
card(0.6, 4.4, 6.0, 2.8,
    "难点三：MQTT 通信 — 阻塞重连 + Broker IP 硬编码",
    "问题一: PubSubClient.connect() 同步阻塞\n"
    "  • Broker 不可达时 TCP connect 超时数秒 → 主循环完全卡死\n"
    "  • 按键扫描/OLED 渲染/倒计时/看门狗全部停摆\n"
    "  • ESP32-C3 单核 RISC-V，无法 xTaskCreatePinnedToCore\n"
    "  • 当前: 5s 冷却窗口 + 正常连接时 mqttClient.loop() 非阻塞\n\n"
    "问题二: Broker IP 写死 (config.h MQTT_BROKER_IP \"172.21.1.146\")\n"
    "  • PC DHCP 换 IP 后失效 → 每 5s 触发一次阻塞式重连\n"
    "  • 两个问题耦合: IP 不可达 → 频繁重连 → 主循环反复卡顿\n\n"
    "补救: 修 IP 动态配置 (NVS 持久化) 即可大幅减少重连触发\n"
    "根治: 换 AsyncMQTTClient (回调非阻塞) 或 FreeRTOS 任务隔离 MQTT",
    PURPLE)

card(6.9, 4.4, 5.5, 2.8,
    "难点四：Flutter APK 构建失败 → Web 退守",
    "问题: flutter build apk 在国内网络环境下全部失败\n"
    "  • 依赖链: Gradle 发行版→插件→Maven→Android SDK→Kotlin→Dart pub\n"
    "  • 配置了阿里云/腾讯云/Flutter国内镜像 + Clash 代理 (127.0.0.1:7890)\n"
    "  • 但镜像无法覆盖全部环节 (pub.dev/dl.google.com/部分 Kotlin)\n"
    "  • .gradle/ 下 8.6/8.9/9.1.0/8.14 四版本 → 反复调版本兼容性\n\n"
    "退守: Flutter App (25文件~2500行) UI已完整，改 Web SPA 方案\n"
    "  • 浏览器即运行时，无需 Gradle/Android SDK\n"
    "  • Cloudflare Pages 部署 desktoppet.pages.dev (静态前端)\n"
    "  • Cloudflare Workers 部署 desktoppet-search.xxx.workers.dev (搜索代理)\n"
    "  • 国内演示仍用本地 Live Server (Cloudflare 国内访问不稳定)\n\n"
    "未来: PWA 包装 Web = 接近原生体验 | 或用完整 Flutter 镜像链重试",
    GRAY)

# page number
tb(11.5, 7.0, 1.5, 0.4, "17/22", sz=12, c=GRAY, align=PP_ALIGN.RIGHT)

prs.save(path)
print("P17 updated OK")
